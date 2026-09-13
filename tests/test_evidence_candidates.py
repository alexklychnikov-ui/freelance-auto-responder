from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from src.evidence.collector import EvidenceService, collect_candidates
from src.evidence.discovery import ResourceCandidate, discover, extract_url_candidates
from src.evidence.files import extract_text_from_bytes
from src.models import ProjectFull

TZ_3252339 = (
    "ЗдравствуйтеНужно разработать аналог сайт торги-россии.рф, интересует только "
    "поиск и выдача торгов с документами, без личного кабинета и прочего. "
    "Там парсится отсюда https://torgi.gov.ru"
)


def _project(desc: str, *, url: str = "https://kwork.ru/projects/3252339") -> ProjectFull:
    return ProjectFull(
        platform="kwork",
        source_key="kwork_dev_it",
        project_id="3252339",
        url=url,
        title="Разработать сайт",
        full_description=desc,
    )


def test_url_trailing_punctuation_stripped() -> None:
    text = "Смотри источник https://torgi.gov.ru/new/public)."
    found = extract_url_candidates(text, max_urls=3)
    assert len(found) == 1
    assert found[0].input_ref == "https://torgi.gov.ru/new/public"
    assert found[0].role == "data_source"


def test_cyrillic_idn_domain() -> None:
    found = extract_url_candidates("аналог торги-россии.рф", max_urls=3)
    assert len(found) == 1
    assert "торги-россии.рф" in found[0].input_ref
    assert found[0].role == "reference"


def test_duplicate_urls_deduped() -> None:
    text = "https://torgi.gov.ru и ещё раз https://torgi.gov.ru/ и www.torgi.gov.ru"
    found = extract_url_candidates(text, max_urls=5)
    assert len(found) == 1


def test_filename_like_hosts_skipped() -> None:
    text = "смотри report.txt data.csv README.md photo.png archive.zip dump.sql"
    found = extract_url_candidates(text, max_urls=10)
    assert found == []


def test_repo_hint_not_triggered_by_report() -> None:
    text = "аналог https://example.com/path смотри report.pdf вложение"
    found = extract_url_candidates(text, max_urls=3)
    assert len(found) == 1
    assert found[0].role == "reference"


def test_3252339_description_candidates() -> None:
    cands = discover(_project(TZ_3252339), max_urls=3)
    urls = [c for c in cands if c.kind == "url"]
    refs = {c.input_ref for c in urls}
    assert any("торги-россии.рф" in r for r in refs)
    assert any("torgi.gov.ru" in r for r in refs)

    by_host = {c.title_hint: c for c in urls}
    # roles
    torgi_rossii = next(c for c in urls if "торги-россии" in c.input_ref)
    torgi_gov = next(c for c in urls if "torgi.gov.ru" in c.input_ref)
    assert torgi_rossii.role == "reference"
    assert torgi_gov.role == "data_source"
    assert isinstance(torgi_rossii, ResourceCandidate)


def test_skip_kwork_project_url() -> None:
    desc = (
        "Ссылка на заказ https://kwork.ru/projects/3252339/view и источник "
        "https://example.com/api/docs"
    )
    cands = collect_candidates(_project(desc), max_urls=3)
    urls = [c.input_ref for c in cands if c.kind == "url"]
    assert urls
    assert all("kwork.ru" not in u for u in urls)
    assert any("example.com" in u for u in urls)


def test_attachment_marker_parsing() -> None:
    desc = (
        "ТЗ в файле\n\n"
        "--- Вложение: Техническое задание.pdf ---\n"
        "текст вложения\n\n"
        "--- Вложение: data.xlsx ---\n"
        "a\tb"
    )
    cands = discover(_project(desc), max_urls=3)
    atts = [c for c in cands if c.kind == "attachment"]
    assert [c.input_ref for c in atts] == [
        "Техническое задание.pdf",
        "data.xlsx",
    ]


def test_xlsx_multi_sheet_extract_smoke() -> None:
    wb = Workbook()
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Lots"
    ws1["A1"] = "lot_id"
    ws1["B1"] = "price"
    ws1["A2"] = 184729
    ws1["B2"] = 7_000_000
    ws2 = wb.create_sheet("Meta")
    ws2["A1"] = "source"
    ws2["B1"] = "torgi.gov.ru"
    buf = BytesIO()
    wb.save(buf)
    text = extract_text_from_bytes(buf.getvalue(), "tz.xlsx")
    assert "[sheet:Lots]" in text
    assert "184729" in text
    assert "[sheet:Meta]" in text
    assert "torgi.gov.ru" in text


def test_xlsx_sheet_and_row_caps() -> None:
    from src.evidence.files import MAX_XLSX_ROWS, MAX_XLSX_SHEETS

    wb = Workbook()
    wb.active.title = "S0"
    for i in range(1, MAX_XLSX_SHEETS + 2):
        wb.create_sheet(f"S{i}")
    for name in wb.sheetnames:
        ws = wb[name]
        for r in range(1, MAX_XLSX_ROWS + 30):
            ws.cell(r, 1, f"{name}-{r}")
    buf = BytesIO()
    wb.save(buf)
    text = extract_text_from_bytes(buf.getvalue(), "caps.xlsx")
    sheets = [line for line in text.splitlines() if line.startswith("[sheet:")]
    assert len(sheets) == MAX_XLSX_SHEETS
    assert f"S0-{MAX_XLSX_ROWS}" in text
    assert f"S0-{MAX_XLSX_ROWS + 1}" not in text
    assert "[sheet:S5]" not in text


def test_oversized_file_rejected() -> None:
    from src.evidence.files import MAX_FILE_BYTES

    blob = b"x" * (MAX_FILE_BYTES + 1)
    assert extract_text_from_bytes(blob, "huge.txt") == ""


def test_unsupported_png_empty() -> None:
    assert extract_text_from_bytes(b"\x89PNG\r\n\x1a\n\x00\x01", "photo.png") == ""


def test_evidence_service_stub() -> None:
    svc = EvidenceService()
    assert svc.collect_candidates(_project("нет ссылок")) == []
    out = svc.collect_candidates(_project(TZ_3252339), max_urls=3)
    assert len(out) >= 2
