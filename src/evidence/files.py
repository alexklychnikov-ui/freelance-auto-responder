from __future__ import annotations

import json
import logging
import re
from io import BytesIO
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from src.adapters.kwork_attachments import extract_text_from_bytes as _base_extract

logger = logging.getLogger(__name__)

MAX_XLSX_SHEETS = 5
MAX_XLSX_ROWS = 200
MAX_XLSX_COLS = 64
MAX_XLSX_CELLS = 2_000
MAX_EXTRACT_CHARS = 20_000
MAX_FILE_BYTES = 5_000_000


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_xlsx(data: bytes) -> str:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    cells_used = 0
    try:
        for sheet_idx, name in enumerate(wb.sheetnames):
            if sheet_idx >= MAX_XLSX_SHEETS:
                break
            ws = wb[name]
            parts.append(f"[sheet:{name}]")
            row_n = 0
            for row in ws.iter_rows(values_only=True):
                row_n += 1
                if row_n > MAX_XLSX_ROWS:
                    break
                vals: list[str] = []
                for col_i, cell in enumerate(row):
                    if col_i >= MAX_XLSX_COLS:
                        break
                    if cells_used >= MAX_XLSX_CELLS:
                        break
                    if cell is None:
                        continue
                    text = str(cell).strip()
                    if not text:
                        continue
                    vals.append(text)
                    cells_used += 1
                if vals:
                    parts.append("\t".join(vals))
                if cells_used >= MAX_XLSX_CELLS:
                    break
            if cells_used >= MAX_XLSX_CELLS:
                break
    finally:
        wb.close()
    return "\n".join(parts).strip()


def _extract_json(data: bytes) -> str:
    raw = _decode_text(data).strip()
    if not raw:
        return ""
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return raw[:MAX_EXTRACT_CHARS]
    try:
        return json.dumps(parsed, ensure_ascii=False, indent=2)[:MAX_EXTRACT_CHARS]
    except (TypeError, ValueError):
        return raw[:MAX_EXTRACT_CHARS]


def _extract_xml(data: bytes) -> str:
    raw = _decode_text(data)
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return re.sub(r"<[^>]+>", " ", raw).strip()[:MAX_EXTRACT_CHARS]
    texts: list[str] = []
    for el in root.iter():
        if el.text and el.text.strip():
            texts.append(el.text.strip())
        if el.tail and el.tail.strip():
            texts.append(el.tail.strip())
    return "\n".join(texts).strip()[:MAX_EXTRACT_CHARS]


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    """Extract plain text from attachment bytes. Unsupported → \"\" (no raise)."""
    if not data:
        return ""
    if len(data) > MAX_FILE_BYTES:
        logger.warning(
            "evidence_file_too_large filename=%s size=%s",
            filename,
            len(data),
        )
        return ""
    name = (filename or "").lower().split("?")[0]
    try:
        if name.endswith((".xlsx", ".xlsm")):
            text = _extract_xlsx(data)
        elif name.endswith(".json"):
            text = _extract_json(data)
        elif name.endswith(".xml"):
            text = _extract_xml(data)
        else:
            text = _base_extract(data, filename)
    except Exception:
        logger.exception("evidence_file_extract_failed filename=%s", filename)
        return ""
    if not text:
        return ""
    return text[:MAX_EXTRACT_CHARS]
